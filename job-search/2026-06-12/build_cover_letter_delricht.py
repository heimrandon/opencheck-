#!/usr/bin/env python3
"""Generates the DelRicht Research Healthcare Project Coordinator cover letter (.docx)."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

OUT = os.path.join(os.path.dirname(__file__), "cover_letters")
os.makedirs(OUT, exist_ok=True)

BODY = [
    "Dear DelRicht Research Hiring Team,",
    "I'm applying for the Healthcare Project Coordinator role, and I'll be direct about why it caught my attention: it's the behind-the-scenes, project-driven side of healthcare I've been deliberately moving toward — work where attention to detail, firm deadlines, and clean execution decide whether something succeeds. As a New Orleans local, I'd be on-site and ready to contribute from day one.",
    "I have a master's in clinical psychology and three-plus years across CrescentCare, Ochsner, and the VA — well beyond your bachelor's-plus-one-year requirement. More importantly, my current work mirrors study start-up operations more than the title might suggest:",
]
BULLETS = [
    "Protocol- and deadline-driven work: At CrescentCare I support a SAMHSA-funded grant program — identifying eligible patients, completing required tools within strict timeframes, and keeping documentation and data clean. It's the same accuracy-under-a-clock rhythm a clinical trial demands.",
    "Process ownership: I redesigned our intake workflow for a 40% efficiency gain, taking initiative to find the bottleneck rather than waiting to be told where it was.",
    "Detail and systems: As a psychometrist at Ochsner I administered and scored standardized assessments with precision, and I'm comfortable in EMR systems and medical terminology.",
    "Multi-stakeholder communication: I coordinate daily across medical, psychiatry, and social-work teams — translating between groups and following through until the work is done.",
]
CLOSE = [
    "Your core values land for me. I have a bias toward action (Production), I'm genuinely coachable and energized by a growing team (Humility), and I stay calm and accurate under pressure (Consistency). I'd bring that same ownership to making sure your sites, sponsors, and CROs get dependable follow-through throughout study start-up.",
    "I'd welcome the chance to talk about how I can help your team launch studies smoothly. Thank you for your consideration.",
]

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for m in doc.sections:
    m.top_margin = m.bottom_margin = Pt(54); m.left_margin = m.right_margin = Pt(72)

# Letterhead
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Randon Heim"); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor(0x1F,0x3A,0x5F)
p.paragraph_format.space_after = Pt(0)
c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
c.add_run("New Orleans, Louisiana  •  601-494-9323  •  heimrandon.rh@gmail.com").font.size = Pt(9.5)
c.paragraph_format.space_after = Pt(10)

for line in ["[Date]", "", "DelRicht Research", "New Orleans, LA", ""]:
    pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(2)
    pp.add_run(line)

for para in BODY:
    pp = doc.add_paragraph(para); pp.paragraph_format.space_after = Pt(8)
for b in BULLETS:
    pp = doc.add_paragraph(style="List Bullet"); pp.add_run(b); pp.paragraph_format.space_after = Pt(4)
for para in CLOSE:
    pp = doc.add_paragraph(para); pp.paragraph_format.space_after = Pt(8)

for line in ["Sincerely,", "Randon Heim"]:
    pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(0); pp.add_run(line)

path = os.path.join(OUT, "Randon_Heim_CoverLetter_DelRicht_Healthcare_Project_Coordinator.docx")
doc.save(path)
print("Wrote", path)
