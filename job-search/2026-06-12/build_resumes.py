#!/usr/bin/env python3
"""
Generates ATS-friendly, role-tailored .docx resumes for Randon Heim.
One shared base (real experience) + per-role headline/summary/keywords.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

OUT = os.path.join(os.path.dirname(__file__), "resumes")
os.makedirs(OUT, exist_ok=True)

NAME = "Randon Heim"
CONTACT = "New Orleans, Louisiana  •  601-494-9323  •  heimrandon.rh@gmail.com  •  [LinkedIn URL]"

# ---- Shared experience (real, from current resume; tightened + quantified) ----
EXPERIENCE = [
    {
        "org": "CrescentCare — New Orleans Health Clinic", "loc": "New Orleans, LA",
        "title": "Behavioral Health Specialist / Group Leader", "dates": "04/2024 – Present",
        "bullets": [
            "Coordinate behavioral health care for a SAMHSA-funded population (PIPBHC grant): identify eligible patients, complete required screening tools within grant timeframes, and maintain clean documentation and data integrity.",
            "Redesigned intake and screening workflow, improving care-coordination efficiency by 40% and reducing time-to-service for patients.",
            "Conduct comprehensive behavioral health screenings and risk stratification, matching patients to appropriate intervention pathways and aftercare resources.",
            "Partner with a multidisciplinary team (medical, psychiatry, social work) to align treatment plans, track progress, and ensure timely follow-up.",
            "Built a patient progress-tracking protocol that improved treatment-plan adherence and follow-up reliability.",
        ],
    },
    {
        "org": "Ochsner Health — Main Campus", "loc": "New Orleans, LA",
        "title": "Psychometrist", "dates": "07/2023 – 01/2024",
        "bullets": [
            "Administered and scored standardized psychological assessments for 75+ patients with precision, under the supervision of John R. Sawyer, ABPP.",
        ],
    },
    {
        "org": "U.S. Department of Veterans Affairs", "loc": "",
        "title": "Psychology Extern", "dates": "07/2022 – 06/2023",
        "bullets": [
            "Coordinated and documented treatment plans for 100+ patients across an outpatient PTSD clinic and an inpatient Mental Health Unit.",
            "Facilitated psychoeducational groups and delivered evidence-based interventions (Prolonged Exposure, Cognitive Processing Therapy, Imagery Rehearsal Therapy).",
        ],
    },
]

EDUCATION = [
    ("M.A., Clinical Psychology — GPA 3.8", "The Chicago School of Professional Psychology", "08/2020 – 05/2023"),
    ("B.S., Neuroscience & Cognitive Studies", "Millsaps College", "2015 – 2019"),
]

PROJECTS = [
    ("Community Outreach Lecture Series — Magnolia Medical Health Foundation (Gulfport, MS)", "06/2022 – 11/2023",
     "Co-led community lectures on diabetes prevention, vaccine hesitancy, work-life balance, and community support for a nonprofit health organization."),
]

# ---- Per-role tailoring ----
ROLES = {
    "Two_Chairs_Care_Coordinator": {
        "headline": "Care Coordinator — Client Journey, Scheduling & Care Operations",
        "summary": ("Behavioral health professional with integrated-care and SAMHSA population-health experience supporting "
                    "clients through every step of their care journey. Skilled in scheduling logistics, answering care, "
                    "billing, and insurance questions, and partnering with clinical teams on the operational workflows that "
                    "keep care running smoothly. Translate clinical practice into reliable, repeatable processes that improve "
                    "the client experience in virtual care settings."),
        "skills": ["Client Care Coordination", "Scheduling & Logistics",
                   "Billing & Insurance Navigation", "Clinical-Team Operations Support",
                   "Behavioral Health Screening & Triage", "Documentation & Data Integrity",
                   "Virtual / Telehealth Care", "Process Improvement"],
    },
    "Included_Health_Member_Care_Advocate": {
        "headline": "Healthcare Member Advocate — Care Navigation & Patient Support",
        "summary": ("Healthcare professional who guides patients through the complexity of care and benefits with empathy and "
                    "accuracy. Experienced in high-volume, patient-facing communication, behavioral health screening, and "
                    "coordinating resources across multidisciplinary teams. Known for first-contact problem solving, clean "
                    "documentation, and a calm, solutions-oriented approach to every interaction."),
        "skills": ["Patient / Member Advocacy", "Healthcare & Benefits Navigation",
                   "High-Volume Member Communication", "Empathetic Service & De-escalation",
                   "Resource Coordination", "First-Contact Resolution",
                   "HIPAA & Documentation", "CRM / Healthcare Software"],
    },
    "Lyra_Health_Client_Support_Specialist": {
        "headline": "Client Support Specialist — Mental Health Care Access & Coordination",
        "summary": ("Behavioral health professional skilled at being the first point of contact for clients seeking mental "
                    "health care. Experienced in intake screening, care coordination, and connecting people to evidence-based "
                    "services with warmth and precision. Strong communicator who resolves issues, documents accurately, and "
                    "collaborates cross-functionally to deliver an excellent client experience."),
        "skills": ["Client Support & Intake", "Mental Health Care Access",
                   "Empathetic Communication", "Care Coordination",
                   "Problem Resolution", "Scheduling & Follow-up",
                   "Documentation Accuracy", "Cross-functional Collaboration"],
    },
}

def set_base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

def heading(doc, text):
    p = doc.add_paragraph()
    p.space_before = Pt(8)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1F3A5F')
    pbdr.append(bottom); pPr.append(pbdr)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p

def build(role_key, cfg):
    doc = Document()
    set_base_style(doc)
    for m in doc.sections:
        m.top_margin = m.bottom_margin = Pt(36)
        m.left_margin = m.right_margin = Pt(54)

    # Header
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(NAME); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(0x1F,0x3A,0x5F)
    p.paragraph_format.space_after = Pt(0)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p2.add_run(cfg["headline"]); rr.italic = True; rr.font.size = Pt(11)
    p2.paragraph_format.space_after = Pt(2)
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run(CONTACT).font.size = Pt(9.5)
    p3.paragraph_format.space_after = Pt(4)

    # Summary
    heading(doc, "Professional Summary")
    doc.add_paragraph(cfg["summary"])

    # Core competencies
    heading(doc, "Core Competencies")
    doc.add_paragraph("  •  ".join(cfg["skills"]))

    # Experience
    heading(doc, "Professional Experience")
    for job in EXPERIENCE:
        ph = doc.add_paragraph(); ph.paragraph_format.space_after = Pt(0)
        loc = f" — {job['loc']}" if job['loc'] else ""
        rr = ph.add_run(f"{job['org']}{loc}"); rr.bold = True
        rr2 = ph.add_run(f"\t{job['dates']}")
        # title line
        tp = doc.add_paragraph(); tp.paragraph_format.space_after = Pt(1)
        ti = tp.add_run(job["title"]); ti.italic = True
        for b in job["bullets"]:
            bullet(doc, b)

    # Education
    heading(doc, "Education")
    for deg, school, dates in EDUCATION:
        ep = doc.add_paragraph(); ep.paragraph_format.space_after = Pt(0)
        ep.add_run(deg).bold = True
        ep.add_run(f"  |  {school}  |  {dates}")

    # Projects
    heading(doc, "Community & Projects")
    for title, dates, desc in PROJECTS:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(0)
        pp.add_run(title).bold = True
        bullet(doc, desc)

    path = os.path.join(OUT, f"Randon_Heim_Resume_{role_key}.docx")
    doc.save(path)
    return path

if __name__ == "__main__":
    for k, c in ROLES.items():
        print("Wrote", build(k, c))
